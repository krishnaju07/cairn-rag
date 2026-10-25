import csv
from pathlib import Path

import pymupdf
from docx import Document

OUT = Path(__file__).parent / "docs"

HANDBOOK = [
    (
        "Welcome to Northwind Labs",
        [
            "Northwind Labs builds measurement software for field researchers. This handbook explains how we work, what you can expect from us and what we expect from you. It applies to every full-time and part-time employee in every office and remote location.",
            "If something in this handbook conflicts with local employment law, the law wins. Questions about any policy can be sent to the People team at people@northwind.example.",
        ],
    ),
    (
        "Working hours",
        [
            "The standard working week is 40 hours. Core hours are 10:00 to 16:00 in your local time zone; you are expected to be reachable during core hours and may arrange the remaining hours around them.",
            "Overtime is not expected. If a project requires extra hours, your manager must approve them in advance, and the time is returned as paid time off within the same quarter.",
        ],
    ),
    (
        "Vacation",
        [
            "New full-time employees get 15 paid vacation days a year, earned at 1.25 days per month from their start date. Part-time employees accrue vacation in proportion to their contracted hours.",
            "After three full years of service the allowance rises to 20 days, and up to 5 unused days can be carried over into the next year. Carried-over days expire on 31 March.",
            "Vacation requests of more than five consecutive days should be submitted at least four weeks ahead so that teams can plan coverage.",
        ],
    ),
    (
        "Sick leave and family leave",
        [
            "Every employee receives 10 paid sick days per year. A doctor's note is required once an absence passes 3 consecutive days.",
            "Primary caregivers receive 16 weeks of fully paid parental leave. Secondary caregivers receive 6 weeks. Leave can be taken at any point during the first 12 months after a child joins the family.",
            "Bereavement leave is 5 paid days for an immediate family member and 2 paid days for extended family.",
        ],
    ),
    (
        "Travel and expenses",
        [
            "Business expenses are reimbursed within 30 days of submission through the expense portal. Receipts are required for any single expense above $25.",
            "Meals during business travel are capped at $75 per day. Hotels are reimbursed up to $220 per night. Flights must be booked in economy class unless the flight is longer than eight hours.",
            "Personal expenses, minibar charges and entertainment are not reimbursable.",
        ],
    ),
    (
        "Equipment and security",
        [
            "Laptops are refreshed every 3 years. A lost or stolen device must be reported to the IT desk within 24 hours so that it can be wiped remotely.",
            "Two-factor authentication is mandatory on all company accounts. Passwords must be stored in the company password manager and never shared over chat or email.",
        ],
    ),
    (
        "Performance reviews",
        [
            "Performance reviews take place twice per year, in June and December. Each review combines a self-assessment, feedback from at least three colleagues and a conversation with your manager.",
            "Salary adjustments are decided after the December review and take effect on 1 February.",
        ],
    ),
]

BENEFITS_SECTIONS = [
    (
        "Health insurance",
        "Northwind pays 80% of the monthly premium for employees and 60% for dependents. Enrollment must be completed within 30 days of your start date; after that you can only change plans during the annual open enrollment in November.",
    ),
    (
        "Retirement savings",
        "The company matches 401(k) contributions up to 4% of salary. Matching contributions vest immediately.",
    ),
    (
        "Learning budget",
        "Each employee has a learning budget of $1,200 per calendar year for courses, books and conferences. Unused budget does not roll over.",
    ),
    (
        "Wellness",
        "A wellness stipend of $50 per month can be used for gym memberships, fitness apps or mental health services. Travel insurance for personal vacations is not covered.",
    ),
]

BENEFITS_PLANS = [
    ("Plan", "Employee monthly cost", "Dependents covered"),
    ("Basic", "$0", "No"),
    ("Plus", "$45", "Spouse only"),
    ("Family", "$190", "Spouse and children"),
]

REMOTE_POLICY = """# Remote Work Policy

Most roles at Northwind Labs can be done remotely. Each team keeps one anchor day, which is Tuesday, when everyone is expected to be in the office or on a video call together.

## Home office

New remote employees receive a one-time home-office stipend of $600. Internet costs are reimbursed at $40 per month on submission of a bill.

## Working from another country

Employees may work from another country for up to 20 days per calendar year with their manager's approval. Longer stays need approval from the People team because of tax and immigration rules.

## Availability

Remote employees follow the same core hours as everyone else, 10:00 to 16:00 local time, and keep their calendar up to date.
"""

HOLIDAYS = [
    ("2026-01-01", "New Year's Day", "Closed"),
    ("2026-01-19", "Martin Luther King Jr. Day", "Closed"),
    ("2026-03-13", "Company Offsite", "Open, in person"),
    ("2026-05-25", "Memorial Day", "Closed"),
    ("2026-07-03", "Independence Day (observed)", "Closed"),
    ("2026-09-07", "Labor Day", "Closed"),
    ("2026-11-26", "Thanksgiving", "Closed"),
    ("2026-11-27", "Day after Thanksgiving", "Closed"),
    ("2026-12-24", "Christmas Eve", "Half day"),
    ("2026-12-25", "Christmas Day", "Closed"),
    ("2026-12-31", "New Year's Eve", "Half day"),
]


def build_handbook() -> None:
    document = pymupdf.open()
    for index, (title, paragraphs) in enumerate(HANDBOOK, start=1):
        page = document.new_page()
        page.insert_textbox(pymupdf.Rect(56, 56, 540, 90), title, fontsize=20, fontname="hebo")
        page.insert_textbox(
            pymupdf.Rect(56, 110, 540, 760), "\n\n".join(paragraphs), fontsize=11.5, fontname="helv", lineheight=1.4
        )
        footer = f"Northwind Labs Employee Handbook 2026 - page {index}"
        page.insert_textbox(pymupdf.Rect(56, 800, 540, 820), footer, fontsize=8, fontname="helv")
    document.set_metadata({"title": "Employee Handbook 2026", "author": "Northwind Labs People Team"})
    document.save(OUT / "Employee Handbook 2026.pdf")


def build_benefits() -> None:
    document = Document()
    document.core_properties.title = "Benefits Guide"
    document.core_properties.author = "Northwind Labs People Team"
    document.add_heading("Benefits Guide", level=0)
    for title, body in BENEFITS_SECTIONS:
        document.add_heading(title, level=1)
        document.add_paragraph(body)
    document.add_heading("Health plans at a glance", level=1)
    table = document.add_table(rows=len(BENEFITS_PLANS), cols=3)
    for row, values in zip(table.rows, BENEFITS_PLANS):
        for cell, value in zip(row.cells, values):
            cell.text = value
    document.save(OUT / "Benefits Guide.docx")


def build_holidays() -> None:
    with open(OUT / "Holiday Calendar 2026.csv", "w", newline="") as handle:
        writer = csv.writer(handle)
        writer.writerow(["Date", "Holiday", "Office status"])
        writer.writerows(HOLIDAYS)


if __name__ == "__main__":
    OUT.mkdir(exist_ok=True)
    build_handbook()
    build_benefits()
    (OUT / "Remote Work Policy.md").write_text(REMOTE_POLICY)
    build_holidays()
    print(f"Wrote corpus to {OUT}")
